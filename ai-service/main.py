"""PhilFreela AI service.

Online, it runs on Vercel (free), which finds `app` in this file by itself;
the live website reaches it through its own /ai address (a Cloudflare
function forwards it, see client/functions/ai). See README.md.

To run it on a laptop instead, first time only, from this folder:
    py -3.10 -m venv .venv
    .venv\\Scripts\\python -m pip install -r requirements.txt
    .venv\\Scripts\\python get_models.py
    copy .env.example .env      (then fill in the real values)

Run it:
    .venv\\Scripts\\python -m uvicorn main:app --port 8000
"""

import io
import logging
import os
import tempfile
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

import cv2
from docx import Document as DocxDocument
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from PIL import Image
from postgrest.exceptions import APIError
from pydantic import BaseModel
from pypdf import PdfReader
from supabase import ClientOptions, create_client

from document_pages import render_pages, watermark_page
from face_check import NoFaceError, check_faces, find_duplicate, prepare_image
from id_qr import check_philsys_qr
from hidden_watermark import new_code, protect_photo, read_code, read_code_from_frames, read_uncropped_codes
from listing_search import search_listings
from moodboard import match_moodboard
from photo_checks import PhotoProblem, check_face_scan, check_id_back, check_id_front
from recommendations import recommend
from similarity import image_embedding
from text_embedder import embed_texts
from text_watermark import add_code as add_text_code
from text_watermark import footer as text_footer
from text_watermark import paragraphs_for_check, strip_hidden
from text_watermark import read_code as read_text_code
from visible_watermark import draw_visible_watermark, watermark_layer
from watermark_video import code_reads_back, output_size, read_video_code, spread, video_info, watermark_video

load_dotenv()

# Errors are written to the terminal running the service, so problems can be
# traced without showing technical details to users.
logger = logging.getLogger("uvicorn.error")

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
if not SUPABASE_URL or not SUPABASE_SERVICE_ROLE_KEY:
    raise RuntimeError("Set SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY (in ai-service/.env on a laptop, or the host's environment variables).")

# The service role key skips the database rules, which is why only this
# service may hold it. That's what lets it save verifications that users
# themselves can't create or edit.
# Photo uploads get 60 seconds instead of the default 20, so a slow internet
# connection doesn't make a verification fail.
supabase = create_client(
    SUPABASE_URL,
    SUPABASE_SERVICE_ROLE_KEY,
    options=ClientOptions(storage_client_timeout=60),
)

BUCKET = "verification-docs"
ID_TYPES = {"philsys", "drivers_license", "passport", "umid", "prc"}
ALLOWED_FORMATS = {"JPEG", "PNG", "WEBP"}
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB, same limit as the storage bucket

# Saved photos are shrunk to at most this many pixels. The ID text stays sharp
# enough for an admin to read, and uploads are about half the size of the
# full photos the face check looks at, which matters on slow internet.
STORED_MAX_SIDE = 1280
UPLOAD_TRIES = 3

app = FastAPI(title="PhilFreela AI Service")

# Website addresses allowed to call this service directly from a browser
# (CORS). The website normally goes through its own /ai address instead (the
# dev server on the laptop, a Cloudflare function on the live site), which
# counts as the same site, so this only matters for direct calls.
allowed_origins = [origin.strip() for origin in os.getenv("ALLOWED_ORIGINS", "http://localhost:5173").split(",") if origin.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type", "Authorization"],
)


@app.get("/health")
def health():
    """Quick check that the service is running."""
    return {"status": "ok"}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_user_id(authorization):
    """Asks Supabase whether the user's login token is real; returns their id."""
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Please log in first.")
    try:
        response = supabase.auth.get_user(authorization.removeprefix("Bearer "))
    except Exception:
        raise HTTPException(401, "Your login has expired. Please log in again.")
    if not response or not response.user:
        raise HTTPException(401, "Your login has expired. Please log in again.")
    return response.user.id


def ensure_can_verify(user_id):
    """Only regular users (with a profile) who aren't already verified or waiting."""
    profile = supabase.table("profiles").select("id").eq("id", user_id).limit(1).execute().data
    if not profile:
        raise HTTPException(403, "Only freelancer and client accounts can verify their identity.")

    active = (
        supabase.table("identity_verifications")
        .select("status")
        .eq("user_id", user_id)
        .in_("status", ["pending", "approved"])
        .limit(1)
        .execute()
        .data
    )
    if active:
        if active[0]["status"] == "approved":
            raise HTTPException(409, "You are already verified.")
        raise HTTPException(409, "Your verification is already waiting for review.")


def read_image(upload, photo_name):
    """Checks an uploaded file is a real JPG/PNG/WEBP under 5 MB and opens it."""
    data = upload.file.read(MAX_FILE_SIZE + 1)
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(400, f"The {photo_name} is too big (max 5 MB).")

    # Opening it with Pillow checks the actual file contents, so a text file
    # renamed to ".jpg" is caught here.
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
    except Exception:
        raise HTTPException(400, f"The {photo_name} is not a valid image.")
    if image.format not in ALLOWED_FORMATS:
        raise HTTPException(400, f"The {photo_name} must be a JPG, PNG, or WEBP image.")

    return prepare_image(image)


def to_clean_jpeg(image, max_side=STORED_MAX_SIDE, quality=85):
    """Saves a smaller copy as a fresh JPEG. Nothing hidden (like the GPS
    location phones store in photos) is copied over."""
    stored = image.copy()
    stored.thumbnail((max_side, max_side))
    buffer = io.BytesIO()
    stored.save(buffer, format="JPEG", quality=quality)
    return buffer.getvalue()


def run_photo_check(check, *images):
    """Runs one of the photo_checks and turns a problem into a "retake" reply."""
    try:
        check(*images)
    except PhotoProblem as problem:
        raise HTTPException(422, str(problem))


def get_caller_user_id(authorization, token):
    """Whose photos these are: the phone's QR token if one is sent, otherwise
    the logged-in user's token."""
    if token:
        link = find_valid_link(token)
        if not link:
            raise HTTPException(410, "This link has expired or was already used. Make a new QR code on your computer.")
        return link["user_id"]
    return get_user_id(authorization)


def handle_submission(user_id, id_type, id_photo, id_back, selfie, selfie_left, selfie_right, selfie_left_half=None, selfie_right_half=None):
    """Shared by the computer and phone (QR) uploads. Re-runs every check
    (the website runs them step by step too, but it can't be trusted to),
    compares the faces, saves the photos, and adds a "pending" verification
    for an admin. The two "halfway" face scan frames are only used for the
    same-person check and aren't saved."""
    if id_type not in ID_TYPES:
        raise HTTPException(400, "Please choose a valid ID type.")
    # Both sides of the card are required; passports have no card back.
    if id_type != "passport" and id_back is None:
        raise HTTPException(400, "Please add a photo of the back of your ID.")
    ensure_can_verify(user_id)

    id_image = read_image(id_photo, "photo of the front of your ID")
    id_back_image = read_image(id_back, "photo of the back of your ID") if id_back is not None else None
    selfie_image = read_image(selfie, "face scan")
    left_image = read_image(selfie_left, "face scan")
    right_image = read_image(selfie_right, "face scan")
    has_halves = selfie_left_half is not None and selfie_right_half is not None
    left_half = read_image(selfie_left_half, "face scan") if has_halves else None
    right_half = read_image(selfie_right_half, "face scan") if has_halves else None

    run_photo_check(check_id_front, id_image)
    if id_back_image is not None:
        run_photo_check(check_id_back, id_image, id_back_image)
    run_photo_check(check_face_scan, selfie_image, left_image, right_image, left_half, right_half)

    try:
        result = check_faces(id_image, selfie_image, left_image, right_image, left_half, right_half)
    except NoFaceError as error:
        # Nothing is saved; the user just retakes the photo.
        raise HTTPException(422, str(error))

    # Two warning flags for the admin (never shown to the user):
    # 1. PhilSys IDs: the name in the ID's QR code vs. the profile name.
    qr_status, qr_name = None, None
    if id_type == "philsys":
        try:
            profile = supabase.table("profiles").select("full_name").eq("id", user_id).limit(1).execute().data
            qr_status, qr_name = check_philsys_qr([id_image, id_back_image], (profile[0].get("full_name") if profile else "") or "")
        except Exception:
            logger.exception("Reading the PhilSys QR failed")
    # 2. The same face on another account's pending or approved verification.
    duplicate_of = None
    try:
        others = (
            supabase.table("identity_verifications")
            .select("user_id, face_embedding")
            .neq("user_id", user_id)
            .in_("status", ["pending", "approved"])
            .execute()
            .data
        )
        duplicate_of = find_duplicate(result["embedding"], [(row["user_id"], row["face_embedding"]) for row in others])
    except Exception:
        logger.exception("The duplicate face check failed")

    verification_id = str(uuid.uuid4())
    folder = f"{user_id}/{verification_id}"
    id_path = f"{folder}/id.jpg"
    id_back_path = f"{folder}/id_back.jpg" if id_back_image is not None else None
    selfie_path = f"{folder}/selfie.jpg"
    selfie_left_path = f"{folder}/selfie_left.jpg"
    selfie_right_path = f"{folder}/selfie_right.jpg"

    uploads = [(id_path, id_image), (selfie_path, selfie_image), (selfie_left_path, left_image), (selfie_right_path, right_image)]
    if id_back_path:
        uploads.append((id_back_path, id_back_image))
    saved_paths = [path for path, _ in uploads]

    storage = supabase.storage.from_(BUCKET)

    def upload(item):
        path, image = item
        data = to_clean_jpeg(image)
        # On slow internet a connection sometimes drops mid-upload, so each
        # photo gets up to 3 tries. "upsert" lets a retry replace a
        # half-finished upload of the same file.
        for attempt in range(UPLOAD_TRIES):
            try:
                storage.upload(path, data, {"content-type": "image/jpeg", "upsert": "true"})
                return
            except Exception:
                if attempt == UPLOAD_TRIES - 1:
                    raise
                logger.warning("Upload of %s failed, trying again", path)
                time.sleep(attempt + 1)

    try:
        # Up to 3 photos upload at the same time instead of one after another.
        with ThreadPoolExecutor(max_workers=3) as pool:
            list(pool.map(upload, uploads))
    except Exception:
        logger.exception("Uploading verification photos failed")
        # e.g. the internet connection dropped: remove whatever did upload,
        # so no half-finished verification is left behind.
        try:
            storage.remove(saved_paths)
        except Exception:
            pass
        raise HTTPException(503, "Couldn't save your photos because of a connection problem. Please try again.")

    try:
        supabase.table("identity_verifications").insert({
            "id": verification_id,
            "user_id": user_id,
            "id_type": id_type,
            "id_photo_path": id_path,
            "id_back_path": id_back_path,
            "selfie_path": selfie_path,
            "selfie_left_path": selfie_left_path,
            "selfie_right_path": selfie_right_path,
            "face_match": result["match"],
            "face_distance": result["distance"],
            # The head turns were already checked above; this adds that all
            # three scan frames show the same person.
            "liveness_passed": result["same_person"],
            "id_qr_status": qr_status,
            "id_qr_name": qr_name,
            # The face's 128 numbers (not a photo), kept only for the
            # duplicate check; deleted with the verification.
            "face_embedding": result["embedding"],
            "duplicate_of": duplicate_of,
        }).execute()
    except APIError as error:
        # Don't leave photos behind for a verification that wasn't saved.
        storage.remove(saved_paths)
        # 23505 = the database's "one pending/approved per user" rule, e.g.
        # the user submitted from the computer and the phone at the same time.
        if error.code == "23505":
            raise HTTPException(409, "Your verification is already waiting for review.")
        raise

    # The AI results stay out of the reply: only admins see them, so nobody
    # can keep retrying photos until they fool the face check.
    return {"id": verification_id, "status": "pending"}


def find_valid_link(token):
    """Returns the QR link if it exists, isn't used, and hasn't expired."""
    try:
        uuid.UUID(token)
    except ValueError:
        return None

    now = datetime.now(timezone.utc).isoformat()
    rows = (
        supabase.table("verification_links")
        .select("token, user_id")
        .eq("token", token)
        .is_("used_at", "null")
        .gt("expires_at", now)
        .limit(1)
        .execute()
        .data
    )
    return rows[0] if rows else None


# ---------------------------------------------------------------------------
# Identity verification endpoints
# ---------------------------------------------------------------------------

@app.post("/checks/id-front")
def check_front(
    photo: UploadFile = File(...),
    token: str | None = Form(default=None),
    authorization: str | None = Header(default=None),
):
    """Instant check of the front of the ID, before the user can go on.
    Nothing is saved. Needs a login or a valid QR token, so strangers can't use it."""
    get_caller_user_id(authorization, token)
    run_photo_check(check_id_front, read_image(photo, "photo of the front of your ID"))
    return {"ok": True}


@app.post("/checks/id-back")
def check_back(
    front: UploadFile = File(...),
    back: UploadFile = File(...),
    token: str | None = Form(default=None),
    authorization: str | None = Header(default=None),
):
    """Instant check of the back of the ID (the front is sent too, to make sure
    it isn't the same photo). Nothing is saved."""
    get_caller_user_id(authorization, token)
    front_image = read_image(front, "photo of the front of your ID")
    run_photo_check(check_id_back, front_image, read_image(back, "photo of the back of your ID"))
    return {"ok": True}


@app.post("/verifications")
def submit_verification(
    id_type: str = Form(...),
    id_photo: UploadFile = File(...),
    selfie: UploadFile = File(...),
    selfie_left: UploadFile = File(...),
    selfie_right: UploadFile = File(...),
    id_back: UploadFile | None = File(default=None),
    selfie_left_half: UploadFile | None = File(default=None),
    selfie_right_half: UploadFile | None = File(default=None),
    authorization: str | None = Header(default=None),
):
    """A logged-in user sends their ID photos and face scan."""
    user_id = get_user_id(authorization)
    return handle_submission(user_id, id_type, id_photo, id_back, selfie, selfie_left, selfie_right, selfie_left_half, selfie_right_half)


@app.post("/phone-links")
def create_phone_link(authorization: str | None = Header(default=None)):
    """A logged-in user without a webcam asks for a QR code link."""
    user_id = get_user_id(authorization)
    ensure_can_verify(user_id)

    # Only the newest QR code works: older unused ones are deleted.
    supabase.table("verification_links").delete().eq("user_id", user_id).is_("used_at", "null").execute()
    link = supabase.table("verification_links").insert({"user_id": user_id}).execute().data[0]
    return {"token": link["token"], "expires_at": link["expires_at"]}


@app.get("/phone-links/{token}")
def check_phone_link(token: str):
    """The phone asks if its QR link still works. Only yes/no, no personal info."""
    return {"valid": find_valid_link(token) is not None}


@app.post("/phone-links/{token}/submit")
def submit_from_phone(
    token: str,
    id_type: str = Form(...),
    id_photo: UploadFile = File(...),
    selfie: UploadFile = File(...),
    selfie_left: UploadFile = File(...),
    selfie_right: UploadFile = File(...),
    id_back: UploadFile | None = File(default=None),
    selfie_left_half: UploadFile | None = File(default=None),
    selfie_right_half: UploadFile | None = File(default=None),
):
    """The phone sends the photos. The QR token takes the place of a login."""
    link = find_valid_link(token)
    if not link:
        raise HTTPException(410, "This link has expired or was already used. Make a new QR code on your computer.")

    result = handle_submission(link["user_id"], id_type, id_photo, id_back, selfie, selfie_left, selfie_right, selfie_left_half, selfie_right_half)

    # Mark it used only after success, so a rejected photo can be retaken.
    supabase.table("verification_links").update({"used_at": datetime.now(timezone.utc).isoformat()}).eq("token", token).execute()
    return result


# ---------------------------------------------------------------------------
# Slideshows for services and portfolio projects (watermarking system,
# steps 1-2). Each shows up to 10 photos and videos. Every one goes through
# here (browsers can't upload to the "slide-media" bucket themselves), so each
# is checked, and from step 3 on it also gets watermarked here.
# ---------------------------------------------------------------------------

SLIDE_BUCKET = "slide-media"
SLIDE_UPLOADS_BUCKET = "slide-uploads"
# What a slide can belong to: its table, the media_slides column pointing at
# it, and what to call it in messages.
SLIDE_OWNERS = {
    "service": ("services", "service_id", "service"),
    "portfolio": ("portfolio_items", "portfolio_item_id", "project"),
}
MAX_SLIDES_PER_ITEM = 10
MAX_SLIDES_PER_FREELANCER = 50
MAX_VIDEOS_PER_FREELANCER = 5
MAX_VIDEO_SIZE = 50 * 1024 * 1024  # same limit as the storage buckets
MAX_VIDEO_SECONDS = 30
# Saved photos are at most this many pixels wide or tall: sharp in the
# slideshow, but never the freelancer's full-quality original.
SLIDE_MAX_SIDE = 1200


def remove_quietly(bucket, path):
    """Deletes a file. If that fails it only leaves an unused file behind."""
    try:
        supabase.storage.from_(bucket).remove([path])
    except Exception:
        logger.warning("Couldn't delete %s from %s", path, bucket)


def ensure_own_item(user_id, owner, item_id):
    """Slides can only be added by the service's or project's owner, and not
    while suspended or banned."""
    table, _, name = SLIDE_OWNERS[owner]
    try:
        uuid.UUID(item_id)
    except ValueError:
        raise HTTPException(404, f"That {name} no longer exists.")

    rows = supabase.table(table).select("freelancer_id").eq("id", item_id).limit(1).execute().data
    if not rows:
        raise HTTPException(404, f"That {name} no longer exists.")
    if rows[0]["freelancer_id"] != user_id:
        raise HTTPException(403, f"You can only add files to your own {name}s.")

    # Same check the posting rules use: a ban, or a suspension that blocks
    # posting (e.g. for spam) and hasn't ended.
    if supabase.rpc("is_posting_blocked", {"target": user_id}).execute().data:
        raise HTTPException(403, "Your account is suspended from posting, so you can't upload right now.")


def ensure_room_for_one_more(user_id):
    """Each freelancer can have 50 items in all: photos and videos (in
    services and portfolio projects) plus portfolio documents."""
    slides = supabase.table("media_slides").select("id", count="exact").eq("freelancer_id", user_id).limit(1).execute().count or 0
    documents = (
        supabase.table("portfolio_items").select("id", count="exact")
        .eq("freelancer_id", user_id).eq("kind", "document").limit(1).execute().count or 0
    )
    if slides + documents >= MAX_SLIDES_PER_FREELANCER:
        raise HTTPException(409, f"You've reached the limit of {MAX_SLIDES_PER_FREELANCER} photos, videos and documents. Delete something to add more.")


def next_slide_position(user_id, owner, item_id, media_type):
    """Checks the upload limits and returns the first free spot (1-10) in the
    service's or project's slideshow."""
    _, column, name = SLIDE_OWNERS[owner]
    used = {row["position"] for row in supabase.table("media_slides").select("position").eq(column, item_id).execute().data}
    if len(used) >= MAX_SLIDES_PER_ITEM:
        raise HTTPException(409, f"A {name} can have at most {MAX_SLIDES_PER_ITEM} photos, videos and documents.")

    # Limits across all of the freelancer's services and projects together,
    # to save storage space.
    ensure_room_for_one_more(user_id)
    mine = supabase.table("media_slides").select("media_type").eq("freelancer_id", user_id).execute().data
    if media_type == "video" and sum(row["media_type"] == "video" for row in mine) >= MAX_VIDEOS_PER_FREELANCER:
        raise HTTPException(409, f"You can have at most {MAX_VIDEOS_PER_FREELANCER} videos. Delete a service or project with a video to add another.")

    return next(spot for spot in range(1, MAX_SLIDES_PER_ITEM + 1) if spot not in used)


def video_seconds(data, extension):
    """How long the video is, in seconds (None if it can't be played).
    OpenCV reads videos from a file, so it's written to a temporary one first."""
    handle, path = tempfile.mkstemp(suffix=f".{extension}")
    try:
        with os.fdopen(handle, "wb") as file:
            file.write(data)
        video = cv2.VideoCapture(path)
        try:
            if not video.isOpened() or not video.grab():
                return None
            # Step through the frames to the end: the last frame's time is the
            # length. Stops early once it's clearly too long.
            while True:
                seconds = video.get(cv2.CAP_PROP_POS_MSEC) / 1000
                if seconds > MAX_VIDEO_SECONDS + 1 or not video.grab():
                    return seconds
        finally:
            video.release()
    finally:
        os.remove(path)


def read_slide_video(user_id, video_path, max_seconds=MAX_VIDEO_SECONDS):
    """Takes the video the browser put in slide-uploads and checks it's a real
    MP4, MOV or WEBM of at most 50 MB and 30 seconds (any length when
    max_seconds is None). Returns (data, extension)."""
    # Only from the caller's own folder, so nobody can take someone else's upload.
    if not video_path.startswith(f"{user_id}/") or ".." in video_path:
        raise HTTPException(403, "That video upload isn't yours.")

    try:
        data = supabase.storage.from_(SLIDE_UPLOADS_BUCKET).download(video_path)
    except Exception:
        raise HTTPException(404, "The video upload wasn't found. Please try again.")
    finally:
        # The temporary copy is never needed again, whatever happens next.
        remove_quietly(SLIDE_UPLOADS_BUCKET, video_path)

    if len(data) > MAX_VIDEO_SIZE:
        raise HTTPException(400, "That video is too big (max 50 MB).")

    # The first bytes show the real file type, whatever the file is called.
    # (MP4 and iPhone MOV files are the same family: both start with "ftyp".)
    if data[4:8] == b"ftyp":
        extension = "mov" if data[8:12] == b"qt  " else "mp4"
    elif data[:4] == b"\x1a\x45\xdf\xa3":
        extension = "webm"
    else:
        raise HTTPException(400, "Videos must be MP4, MOV or WEBM.")

    seconds = video_seconds(data, extension)
    if seconds is None:
        raise HTTPException(400, "That video couldn't be opened. Please try another one.")
    if max_seconds and seconds > max_seconds + 0.5:
        raise HTTPException(400, f"Videos can be at most {max_seconds} seconds long.")
    return data, extension


def upload_slide_file(path, data, content_type):
    """Saves a slide to slide-media, with up to 3 tries on a bad connection."""
    storage = supabase.storage.from_(SLIDE_BUCKET)
    for attempt in range(UPLOAD_TRIES):
        try:
            storage.upload(path, data, {"content-type": content_type, "upsert": "true"})
            return
        except Exception:
            if attempt == UPLOAD_TRIES - 1:
                logger.exception("Uploading slide %s failed", path)
                raise HTTPException(503, "Couldn't save that file because of a connection problem. Please try again.")
            logger.warning("Upload of %s failed, trying again", path)
            time.sleep(attempt + 1)


def watermark_photo(user_id, image, promo):
    """Step 3: draws the freelancer's visible watermark (in their style from
    Settings > Watermark Settings, and not on promos), then hides a new
    invisible code in the photo (HiDDeN, see hidden_watermark.py).
    Also makes the copy check's numbers (step 5, see similarity.py) for the
    photo as uploaded and, if a visible watermark was drawn, as shown.
    Returns (jpeg_bytes, code, embeddings), where code is None for a picture
    too plain to hold one."""
    image = image.copy()
    image.thumbnail((SLIDE_MAX_SIDE, SLIDE_MAX_SIDE))
    embeddings = {"uploaded": image_embedding(image)}

    shown = image
    if not promo:
        shown = draw_visible_watermark(image, *watermark_style(user_id))
        if shown is not image:
            embeddings["shown"] = image_embedding(shown)

    code = new_code()
    data, has_code = protect_photo(shown, code)
    return data, (code if has_code else None), embeddings


def watermark_style(user_id):
    """(settings, username, full_name) for drawing the freelancer's visible
    watermark (settings is None if they never saved their own style)."""
    settings = (
        supabase.table("watermark_settings")
        .select("visible_enabled, text_mode, custom_text, position, opacity, size, color, show_badge")
        .eq("freelancer_id", user_id).limit(1).execute().data
    )
    profile = supabase.table("profiles").select("username, full_name").eq("id", user_id).limit(1).execute().data
    profile = profile[0] if profile else {}
    return (settings[0] if settings else None), profile.get("username", ""), profile.get("full_name", "")


def page_path(file_path, number):
    """Where a document slide's page picture is saved (step 11): next to its
    text file, "<user id>/<slide id>.txt" -> "<user id>/<slide id>-p1.jpg".
    The website builds the same names (slidePagePaths in lib/slides.js)."""
    return f"{file_path.rsplit('.', 1)[0]}-p{number}.jpg"


def page_pictures(data):
    """Step 11: the first pages of a PDF as pictures (see document_pages.py).
    Returns [] for a DOCX or TXT, and for a PDF whose pages can't be drawn:
    those stay text-only, as before."""
    # The first bytes show the real file type, whatever the file is called.
    if data[:5] != b"%PDF-":
        return []
    try:
        return render_pages(data)
    except Exception:
        logger.warning("Couldn't draw the pages of a PDF; it is saved as text only")
        return []


def watermarked_pages(user_id, pictures, code):
    """Step 11: the page pictures as JPEGs, each with the freelancer's name
    across it and the document's hidden code inside."""
    if not pictures:
        return []
    style = watermark_style(user_id)
    return [watermark_page(picture, code, *style)[0] for picture in pictures]


# Step 7: 5 frames of each video go through the ViT for the copy check.
VIDEO_CHECK_FRAMES = 5


def watermark_video_upload(user_id, data, extension, promo):
    """Step 7: turns the video into a watermarked 720p MP4 (see
    watermark_video.py): the visible watermark on every frame (not on promos)
    and the invisible code in every frame, read back afterwards as a
    self-check. Also makes the copy check's numbers for 5 frames, as uploaded
    and as shown, and reads any hidden code the upload already carries (step
    9). Returns (mp4_bytes, code, embeddings, earlier_code); code is None if
    the finished video didn't give it back."""
    with tempfile.TemporaryDirectory() as folder:
        source, finished = os.path.join(folder, f"upload.{extension}"), os.path.join(folder, "watermarked.mp4")
        with open(source, "wb") as file:
            file.write(data)
        layer = None
        if not promo:
            layer = watermark_layer(output_size(*video_info(source)[:2]), *watermark_style(user_id))
        code = new_code()
        try:
            uploaded, shown = watermark_video(source, finished, code, layer)
        except Exception:
            logger.exception("Watermarking a video failed")
            raise HTTPException(400, "That video couldn't be processed. Please try another MP4, MOV or WEBM video.")
        has_code = code_reads_back(finished, code)
        # Step 9: the code the upload already carries, if any (someone posting
        # a download of a PhilFreela video), read from its key frames as they
        # came in, before our own code went on. Up to 16 of them, as many as
        # Check Ownership reads (step 12; it was 5).
        earlier_code = read_code_from_frames(spread(uploaded, EXTRACT_VIDEO_FRAMES))

        # The copy check's numbers for 5 of the key frames, as they came in,
        # and as shown when they carry a visible watermark.
        versions = [("uploaded", uploaded)] + ([("shown", shown)] if layer is not None else [])
        embeddings = {
            f"{version}_frame{i}": image_embedding(Image.fromarray(frame))
            for version, frames in versions
            for i, frame in enumerate(spread(frames, VIDEO_CHECK_FRAMES), start=1)
        }
        with open(finished, "rb") as file:
            return file.read(), (code if has_code else None), embeddings, earlier_code


# Step 5, the copy check: a new photo whose ViT numbers are at least this
# similar to another freelancer's photo is flagged for an admin. Chosen from
# tests: edited copies (compressed, resized, cropped, screenshots, mirrored)
# scored 0.90-1.00, while different pictures stayed below 0.88 (except the
# same poster template with other words, which an admin can clear).
COPY_CUTOFF = 0.88
# Step 12: a photo at least this similar to another freelancer's also gets
# Check Ownership's 9 "cropped" readings (find_code_owner). Those take about
# 15 times as long as the one reading every photo gets, so photos that look
# nothing like anyone else's skip them. Different pictures scored 0.34 on
# average in the step 5 tests.
FULL_CHECK_FROM = 0.70


def vector_text(numbers):
    """The numbers in the text form the database's vector type reads."""
    return "[" + ",".join(f"{x:.6f}" for x in numbers) + "]"


def copy_check(user_id, embeddings):
    """Compares a new photo (or the frames of a new video) with other
    freelancers' photos and video frames (as uploaded and as shown). Returns
    (status, matched_slide_id, match_score, closest) for the closest one.
    closest: how similar that one is, also when it isn't held (0 when there
    is nothing to compare with)."""
    best = None
    for embedding in embeddings:
        rows = supabase.rpc(
            "closest_slide_embeddings", {"query": vector_text(embedding), "exclude_freelancer": user_id, "how_many": 1}
        ).execute().data
        if rows and (best is None or rows[0]["similarity"] > best["similarity"]):
            best = rows[0]
    closest = best["similarity"] if best else 0
    if closest >= COPY_CUTOFF:
        # (Below 1.0 even for an exact copy: 1.0 means "found by the hidden code".)
        return "flagged", best["slide_id"], min(round(closest, 4), 0.9999), closest
    return "active", None, None, closest


def someone_elses_code(user_id, match):
    """Step 9, the ownership check when posting: does the upload already carry
    another freelancer's hidden code (someone posting a download of their
    photo or video)? match: the saved code the upload's readings matched (from
    find_code_owner or saved_code), or None. Returns that saved code's row
    (with the slide it came from), or None. Re-posting your own work is fine."""
    return match if match and match["freelancer_id"] != user_id else None


@app.post("/slides")
def add_slide(
    service_id: str | None = Form(default=None),
    portfolio_item_id: str | None = Form(default=None),
    image: UploadFile | None = File(default=None),
    video_path: str | None = Form(default=None),
    document: UploadFile | None = File(default=None),
    promo: bool = Form(default=False),
    authorization: str | None = Header(default=None),
):
    """Adds one photo, video or document to the end of the caller's own
    service or portfolio project (send service_id or portfolio_item_id).
    promo: the freelancer marked it as an ad, so it gets no visible watermark.
    Photos and documents come with the request. Videos are too big for that
    (Vercel allows 4.5 MB per request), so the browser uploads them to
    slide-uploads first and sends where it put them.
    A document is saved as its text; a PDF also as pictures of its first
    pages (step 11, page_count in the reply)."""
    user_id = get_user_id(authorization)
    if (service_id is None) == (portfolio_item_id is None):
        raise HTTPException(400, "Send one service or one project.")
    if [image, video_path, document].count(None) != 2:
        raise HTTPException(400, "Send one photo, video or document.")

    owner, item_id = ("service", service_id) if service_id is not None else ("portfolio", portfolio_item_id)
    _, column, name = SLIDE_OWNERS[owner]
    ensure_own_item(user_id, owner, item_id)
    media_type = "image" if image is not None else "video" if video_path is not None else "document"
    position = next_slide_position(user_id, owner, item_id, media_type)

    code, embeddings, text_embeddings, pages = None, {}, [], []
    status, matched_slide_id, matched_item_id, match_score = "active", None, None, None
    if image is not None:
        # Same checks as ID photos (a real JPG/PNG/WEBP under 5 MB), then the
        # watermarks, saved as a fresh JPEG (nothing hidden, like the GPS
        # location phones store in photos, is copied over).
        picture = read_image(image, "photo")
        data, code, embeddings = watermark_photo(user_id, picture, promo)
        extension, content_type = "jpg", "image/jpeg"
        # The ownership check. Is it nearly the same as another freelancer's
        # photo (step 5)? Does it already carry another freelancer's hidden
        # code (step 9)? Either way it waits for an admin; the code is the
        # surer proof, so it wins. The code is read the way Check Ownership
        # reads it (step 12): as the photo is, and, when the photo looks at
        # least a bit like someone else's, also as if it had been cropped.
        status, matched_slide_id, match_score, closest = copy_check(user_id, [embeddings["uploaded"]])
        earlier = someone_elses_code(user_id, find_code_owner(picture, cropped_too=closest >= FULL_CHECK_FROM))
        if earlier:
            status, matched_slide_id, match_score = "flagged", earlier["slide_id"], 1.0
    elif video_path is not None:
        # Checked (MP4, MOV or WEBM, at most 50 MB and 30 seconds), then saved
        # as a watermarked 720p MP4 (step 7).
        raw, raw_extension = read_slide_video(user_id, video_path)
        data, code, embeddings, earlier_code = watermark_video_upload(user_id, raw, raw_extension, promo)
        extension, content_type = "mp4", "video/mp4"
        # The same two checks as photos, on the video's frames.
        earlier = someone_elses_code(user_id, saved_code([earlier_code]))
        if earlier:
            status, matched_slide_id, match_score = "flagged", earlier["slide_id"], 1.0
        else:
            frames_as_uploaded = [e for version, e in embeddings.items() if version.startswith("uploaded")]
            status, matched_slide_id, match_score, _ = copy_check(user_id, frames_as_uploaded)
    else:
        # Writing (step 8): a PDF, DOCX or TXT, handled like portfolio writing.
        # Its text gets the footer and the invisible code, is compared with
        # other freelancers' writing, and is saved as a .txt file.
        raw = read_document(None, document)
        clean = strip_hidden(raw).strip()
        text_embeddings = embed_texts(paragraphs_for_check(clean))
        status, matched_item_id, matched_slide_id, match_score = check_copied_writing(user_id, raw, text_embeddings)
        body, code = watermark_writing(user_id, clean)
        data, extension, content_type = body.encode("utf-8"), "txt", "text/plain"
        # Step 11: a PDF's first pages are also kept as watermarked pictures,
        # so viewers see its real layout. They carry the same code as the text.
        document.file.seek(0)
        pictures = page_pictures(document.file.read())
        # Step 12: a page that already carries another freelancer's hidden
        # code (their photo, or a page of their PDF, saved into this one)
        # holds the document too, like posting it as a photo would.
        earlier = someone_elses_code(user_id, saved_code([read_code(page) for page in pictures])) if pictures else None
        if earlier and match_score != 1.0:
            status, matched_item_id, matched_slide_id, match_score = "flagged", None, earlier["slide_id"], 1.0
        pages = watermarked_pages(user_id, pictures, code)

    slide_id = str(uuid.uuid4())
    file_path = f"{user_id}/{slide_id}.{extension}"
    # Every file of this slide: the slide itself, then a document's page pictures.
    saved_paths = [file_path] + [page_path(file_path, number) for number in range(1, len(pages) + 1)]

    def remove_saved_files():
        """Don't leave files behind for a slide that wasn't saved."""
        for path in saved_paths:
            remove_quietly(SLIDE_BUCKET, path)

    try:
        upload_slide_file(file_path, data, content_type)
        for path, page in zip(saved_paths[1:], pages):
            upload_slide_file(path, page, "image/jpeg")
    except HTTPException:
        remove_saved_files()
        raise

    try:
        supabase.table("media_slides").insert({
            "id": slide_id,
            "freelancer_id": user_id,
            column: item_id,
            "position": position,
            "media_type": media_type,
            "file_path": file_path,
            "page_count": len(pages),
            "watermarked": code is not None,
            "promo": promo,
            "status": status,
            "matched_slide_id": matched_slide_id,
            "matched_item_id": matched_item_id,
            "match_score": match_score,
        }).execute()
    except APIError as error:
        remove_saved_files()
        if error.code == "23505":  # another upload took the same spot at the same moment
            raise HTTPException(409, "Another upload was saving at the same time. Please try again.")
        if error.code == "23503":  # the service or project was deleted meanwhile
            raise HTTPException(404, f"That {name} no longer exists.")
        raise

    # The invisible code goes in its own private table (only this service can read it).
    if code is not None:
        try:
            supabase.table("watermark_codes").insert({"code": code, "slide_id": slide_id, "freelancer_id": user_id}).execute()
        except APIError:
            # Practically impossible (the same random code twice), but then undo the slide.
            logger.exception("Saving the watermark code for slide %s failed", slide_id)
            supabase.table("media_slides").delete().eq("id", slide_id).execute()
            remove_saved_files()
            raise HTTPException(409, "Something went wrong while saving that file. Please try again.")

    # The copy check's numbers, for comparing future uploads with this photo,
    # video or document (private tables). If saving them fails, the slide
    # itself is still fine.
    try:
        if embeddings:
            supabase.table("slide_embeddings").insert([
                {"slide_id": slide_id, "version": version, "freelancer_id": user_id, "embedding": vector_text(numbers)}
                for version, numbers in embeddings.items()
            ]).execute()
        if len(text_embeddings):
            supabase.table("document_embeddings").insert([
                {"slide_id": slide_id, "piece": i, "freelancer_id": user_id, "embedding": vector_text(e)}
                for i, e in enumerate(text_embeddings)
            ]).execute()
    except APIError:
        logger.exception("Saving the copy check numbers for slide %s failed", slide_id)

    # held_because (step 10): why it waits for an admin, so the page can say
    # so: "watermark" (it carries another freelancer's hidden code) or
    # "similar" (it's nearly the same as another freelancer's work).
    held_because = None
    if status == "flagged":
        held_because = "watermark" if match_score == 1.0 else "similar"
    return {
        "id": slide_id, "position": position, "media_type": media_type, "file_path": file_path,
        "page_count": len(pages),
        "watermarked": code is not None, "promo": promo, "status": status, "held_because": held_because,
    }


# ---------------------------------------------------------------------------
# Check Ownership (watermarking system, step 4): the Extraction API.
# Reads the invisible code from a picture someone found and finds the saved
# code closest to it, to say whose work it is.
# ---------------------------------------------------------------------------

# A reading matches a saved code when at most this many of the 48 bits differ.
# (A picture that isn't from PhilFreela gets that close to a given code about
# once in 20 million tries.)
MAX_WRONG_BITS_MATCH = 6
# Several readings of one file (the "un-cropped" readings, the pages of a
# PDF) are extra tries, so they must be a bit closer (more tries = more
# chances of a lucky, wrong match).
MAX_WRONG_BITS_GUESS = 5
# The best match must also be clearly closer than the next closest code.
MIN_GAP_TO_NEXT = 4


def closest_codes(readings):
    """The two saved codes closest to any of the readings, best first
    (the database function closest_watermark_codes, only this service may call it)."""
    return supabase.rpc("closest_watermark_codes", {"candidates": readings, "how_many": 2}).execute().data


def is_match(rows, max_wrong_bits):
    if not rows or rows[0]["wrong_bits"] > max_wrong_bits:
        return False
    next_closest = rows[1]["wrong_bits"] if len(rows) > 1 else 48
    return next_closest - rows[0]["wrong_bits"] >= MIN_GAP_TO_NEXT


def saved_code(readings):
    """The saved code that these readings of one file match (its row from
    closest_codes), or None. One reading may be 6 bits off; several readings
    must be a bit closer (5)."""
    rows = closest_codes(readings)
    limit = MAX_WRONG_BITS_MATCH if len(readings) == 1 else MAX_WRONG_BITS_GUESS
    return rows[0] if is_match(rows, limit) else None


def find_code_owner(picture, cropped_too=True):
    """Whose hidden code does this picture carry? The one function behind
    both the Check Ownership page and the check on every upload (step 12), so
    the two can't drift apart. Returns the saved code's row, or None.
    cropped_too=False skips the second reading, which is the slow one."""
    # 1. The picture as it is.
    match = saved_code([read_code(picture)])
    # 2. As if its edges had been cropped off (see UNCROP_GUESSES in hidden_watermark.py).
    if match is None and cropped_too:
        match = saved_code(read_uncropped_codes(picture))
    return match


def describe_match(match, user_id):
    """What the Check Ownership page shows: the owner, and the service or
    portfolio project the picture came from (unless it was deleted since)."""
    owner_id = match["freelancer_id"]
    owner = supabase.table("profiles").select("id, full_name, username").eq("id", owner_id).limit(1).execute().data
    owner = owner[0] if owner else {"id": owner_id, "full_name": None, "username": None}
    owner["verified"] = bool(supabase.rpc("is_verified", {"target": owner_id}).execute().data)

    source = None
    slide = []
    if match["slide_id"]:
        slide = (
            supabase.table("media_slides").select("file_path, media_type, page_count, service_id, portfolio_item_id")
            .eq("id", match["slide_id"]).limit(1).execute().data
        )
    if slide:
        slide = slide[0]
        kind, table, item_id = (
            ("service", "services", slide["service_id"]) if slide["service_id"]
            else ("project", "portfolio_items", slide["portfolio_item_id"])
        )
        item = supabase.table(table).select("title").eq("id", item_id).limit(1).execute().data
        # The picture shown as "the original". A screenshot of a document's
        # page (step 11) is found by the document's code: show its page 1.
        picture_path = slide["file_path"]
        if slide["media_type"] == "document":
            picture_path = page_path(slide["file_path"], 1) if slide["page_count"] else None
        source = {"kind": kind, "title": item[0]["title"] if item else None, "file_path": picture_path}

    return {
        "found": True,
        "bits_matched": 48 - match["wrong_bits"],
        "is_you": owner_id == user_id,
        "owner": owner,
        "source": source,
        "uploaded_at": match["created_at"],
    }


@app.post("/watermarks/extract")
def extract_watermark(image: UploadFile = File(...), authorization: str | None = Header(default=None)):
    """Check Ownership: whose work is this picture? Needs a login, so strangers
    can't use it. Returns {"found": false} when there's no PhilFreela code."""
    user_id = get_user_id(authorization)
    picture = read_image(image, "picture")
    match = find_code_owner(picture)
    return describe_match(match, user_id) if match else {"found": False}


# Check Ownership for videos reads this many frames together (step 7).
EXTRACT_VIDEO_FRAMES = 16


@app.post("/watermarks/extract-video")
def extract_video_watermark(video_path: str = Form(...), authorization: str | None = Header(default=None)):
    """Check Ownership for videos. The browser puts the video in slide-uploads
    first (it's too big to send directly), then the code is read from 16
    frames at once: every frame carries the same code, so their readings are
    added up bit by bit. (A screenshot of any frame works with the Picture
    check too.)"""
    user_id = get_user_id(authorization)
    data, extension = read_slide_video(user_id, video_path, max_seconds=None)
    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, f"found.{extension}")
        with open(path, "wb") as file:
            file.write(data)
        code = read_video_code(path, EXTRACT_VIDEO_FRAMES)
    if code is None:
        return {"found": False}
    match = saved_code([code])
    return describe_match(match, user_id) if match else {"found": False}


# ---------------------------------------------------------------------------
# Documents (watermarking system, step 6): writing in portfolios gets an
# invisible code (text_watermark.py), a visible footer, and a copy check with
# a text model (text_embedder.py).
# ---------------------------------------------------------------------------

MAX_DOCUMENT_BYTES = 4 * 1024 * 1024  # under Vercel's 4.5 MB per request
MAX_DOCUMENT_CHARACTERS = 20000
MIN_DOCUMENT_WORDS = 10
# A piece of a new document at least this similar in meaning to a piece of
# another freelancer's document counts as copied. Chosen from tests: copies
# with some words changed scored 0.96-0.99; full rewordings 0.59-0.72 and
# different texts on the same topic at most 0.57.
TEXT_COPY_CUTOFF = 0.80


def read_document(text, document):
    """The writing from pasted text or an uploaded TXT, DOCX or PDF file."""
    if text is not None:
        raw = text
    elif document is not None:
        data = document.file.read(MAX_DOCUMENT_BYTES + 1)
        if len(data) > MAX_DOCUMENT_BYTES:
            raise HTTPException(400, "That file is too big (max 4 MB).")
        # The first bytes show the real file type, whatever the file is called.
        try:
            if data[:5] == b"%PDF-":
                pdf = PdfReader(io.BytesIO(data))
                raw = "\n\n".join(page.extract_text() or "" for page in pdf.pages)
            elif data[:2] == b"PK":
                raw = "\n".join(paragraph.text for paragraph in DocxDocument(io.BytesIO(data)).paragraphs)
            else:
                raw = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(400, "Please upload a TXT, DOCX or PDF file (TXT files must be saved as UTF-8).")
        except Exception:
            raise HTTPException(400, "That file couldn't be read. Please upload a TXT, DOCX or PDF file.")
    else:
        raise HTTPException(400, "Paste some text or choose a file.")

    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    words = len(strip_hidden(raw).split())
    if words < MIN_DOCUMENT_WORDS:
        if document is not None and document.filename and document.filename.lower().endswith(".pdf"):
            raise HTTPException(400, "This PDF has no text we can read (it may be a scan). Please use a text PDF, DOCX or TXT.")
        raise HTTPException(400, f"Please add at least {MIN_DOCUMENT_WORDS} words.")
    if len(strip_hidden(raw)) > MAX_DOCUMENT_CHARACTERS:
        raise HTTPException(400, f"That's too long: documents can have up to {MAX_DOCUMENT_CHARACTERS:,} characters.")
    return raw


def closest_documents(piece_embeddings, exclude_freelancer, how_many=1):
    """The documents with a piece closest to any of these pieces (the database
    function closest_document_pieces, only this service may call it)."""
    return supabase.rpc("closest_document_pieces", {
        "queries": [vector_text(e) for e in piece_embeddings],
        "exclude_freelancer": exclude_freelancer,
        "how_many": how_many,
    }).execute().data


def check_copied_writing(user_id, raw, embeddings):
    """The copy check for writing (portfolio documents and, from step 8,
    document slides in services). Returns (status, matched_item_id,
    matched_slide_id, match_score): the portfolio document or document slide
    it matched, if any."""
    # 1. Pasted from another freelancer's writing, hidden code and all.
    found_code = read_text_code(raw)
    if found_code is not None:
        owner = (
            supabase.table("watermark_codes").select("freelancer_id, portfolio_item_id, slide_id")
            .eq("code", found_code).limit(1).execute().data
        )
        if owner and owner[0]["freelancer_id"] != user_id:
            return "flagged", owner[0]["portfolio_item_id"], owner[0]["slide_id"], 1.0
    # 2. Nearly the same meaning as a piece of another freelancer's writing.
    rows = closest_documents(embeddings, user_id)
    if rows and rows[0]["similarity"] >= TEXT_COPY_CUTOFF:
        # (Below 1.0 even for an exact copy: 1.0 means "found by the hidden code".)
        return "flagged", rows[0]["portfolio_item_id"], rows[0]["slide_id"], min(round(rows[0]["similarity"], 4), 0.9999)
    return "active", None, None, None


def watermark_writing(user_id, clean):
    """The writing's watermark: the visible footer (unless turned off), then
    the invisible code after the first word of every sentence. Returns
    (body, code)."""
    settings = (
        supabase.table("watermark_settings").select("text_mode, custom_text, document_footer")
        .eq("freelancer_id", user_id).limit(1).execute().data
    )
    settings = settings[0] if settings else {}
    profile = supabase.table("profiles").select("username, full_name").eq("id", user_id).limit(1).execute().data
    profile = profile[0] if profile else {}
    body = clean
    if settings.get("document_footer", True):
        body += "\n\n" + text_footer(settings, profile.get("username", ""), profile.get("full_name", ""))
    code = new_code()
    return add_text_code(body, code), code


@app.post("/portfolio/documents")
def add_document(
    title: str = Form(...),
    description: str | None = Form(default=None),
    text: str | None = Form(default=None),
    document: UploadFile | None = File(default=None),
    authorization: str | None = Header(default=None),
):
    """Adds writing to the caller's portfolio: pasted text, or a TXT, DOCX or
    PDF file (only the text is kept). It gets the invisible code, the footer,
    and the copy check before it's saved."""
    user_id = get_user_id(authorization)
    title = title.strip()
    if not 1 <= len(title) <= 100:
        raise HTTPException(400, "Please give the document a title (up to 100 characters).")
    description = (description or "").strip() or None
    if description and len(description) > 1000:
        raise HTTPException(400, "The description can be at most 1,000 characters.")

    profile = supabase.table("profiles").select("account_type, username, full_name").eq("id", user_id).limit(1).execute().data
    if not profile or profile[0]["account_type"] != "freelancer":
        raise HTTPException(403, "Only freelancers can add to a portfolio.")
    if supabase.rpc("is_posting_blocked", {"target": user_id}).execute().data:
        raise HTTPException(403, "Your account is suspended from posting, so you can't upload right now.")
    ensure_room_for_one_more(user_id)

    raw = read_document(text, document)
    clean = strip_hidden(raw).strip()
    embeddings = embed_texts(paragraphs_for_check(clean))
    status, matched_item_id, matched_slide_id, match_score = check_copied_writing(user_id, raw, embeddings)
    body, code = watermark_writing(user_id, clean)

    item = supabase.table("portfolio_items").insert({
        "freelancer_id": user_id, "kind": "document", "title": title, "description": description, "body": body,
        "status": status, "matched_item_id": matched_item_id, "matched_slide_id": matched_slide_id,
        "match_score": match_score,
    }).execute().data[0]
    try:
        supabase.table("watermark_codes").insert({"code": code, "portfolio_item_id": item["id"], "freelancer_id": user_id}).execute()
        supabase.table("document_embeddings").insert([
            {"portfolio_item_id": item["id"], "piece": i, "freelancer_id": user_id, "embedding": vector_text(e)}
            for i, e in enumerate(embeddings)
        ]).execute()
    except APIError:
        # Practically impossible (the same random code twice): undo the document.
        logger.exception("Saving the code or copy check numbers for document %s failed", item["id"])
        supabase.table("portfolio_items").delete().eq("id", item["id"]).execute()
        raise HTTPException(409, "Something went wrong while saving that document. Please try again.")

    return {key: item[key] for key in ("id", "freelancer_id", "kind", "title", "description", "body", "status", "created_at")}


def describe_document_match(owner_id, item_id, slide_id, user_id, how, similarity=None, found_at=None):
    """What Check Ownership shows for writing: the owner, and the portfolio
    document or the service (for a document slide, step 8) it came from.
    document is None if that was deleted since."""
    owner = supabase.table("profiles").select("id, full_name, username").eq("id", owner_id).limit(1).execute().data
    owner = owner[0] if owner else {"id": owner_id, "full_name": None, "username": None}
    owner["verified"] = bool(supabase.rpc("is_verified", {"target": owner_id}).execute().data)
    found = None
    if item_id:
        item = supabase.table("portfolio_items").select("id, title, created_at").eq("id", item_id).limit(1).execute().data
        if item:
            found = {**item[0], "kind": "document"}
    elif slide_id:
        slide = supabase.table("media_slides").select("id, service_id, created_at").eq("id", slide_id).limit(1).execute().data
        if slide and slide[0]["service_id"]:
            service = supabase.table("services").select("title").eq("id", slide[0]["service_id"]).limit(1).execute().data
            found = {"id": slide[0]["id"], "title": service[0]["title"] if service else None, "created_at": slide[0]["created_at"], "kind": "service"}
    return {
        "found": True,
        "how": how,  # "code" (the hidden code) or "similarity" (the text model)
        "similarity": similarity,
        "is_you": owner_id == user_id,
        "owner": owner,
        "document": found,
        "uploaded_at": found["created_at"] if found else found_at,
    }


@app.post("/watermarks/extract-text")
def extract_text_watermark(
    text: str | None = Form(default=None),
    document: UploadFile | None = File(default=None),
    authorization: str | None = Header(default=None),
):
    """Check Ownership for writing: whose document is this text from? First
    the hidden code; if it was removed (retyped, pasted as plain text), the
    text model looks for a document with nearly the same meaning."""
    user_id = get_user_id(authorization)
    raw = read_document(text, document)

    code = read_text_code(raw)
    if code is not None:
        row = (
            supabase.table("watermark_codes").select("freelancer_id, portfolio_item_id, slide_id, created_at")
            .eq("code", code).limit(1).execute().data
        )
        if row:
            row = row[0]
            return describe_document_match(row["freelancer_id"], row["portfolio_item_id"], row["slide_id"], user_id, "code", found_at=row["created_at"])

    rows = closest_documents(embed_texts(paragraphs_for_check(strip_hidden(raw))), None)
    if rows and rows[0]["similarity"] >= TEXT_COPY_CUTOFF:
        best = rows[0]
        return describe_document_match(best["freelancer_id"], best["portfolio_item_id"], best["slide_id"], user_id, "similarity", round(best["similarity"], 4))

    return {"found": False}


# ---------------------------------------------------------------------------
# AI search box: the content-based filtering part of the Hybrid
# recommendation system (see listing_search.py)
# ---------------------------------------------------------------------------

class SearchRequest(BaseModel):
    query: str


@app.post("/search")
def search_posts(body: SearchRequest, authorization: str | None = Header(default=None)):
    """The top search box: services and job posts closest in meaning to the
    typed words. Returns ids and scores only; the website loads the posts
    itself, under the normal database rules."""
    get_user_id(authorization)
    query = " ".join(body.query.split())
    if not 2 <= len(query) <= 200:
        raise HTTPException(400, "Type 2 to 200 characters to search.")
    try:
        return {"results": search_listings(supabase, query)}
    except APIError:
        logging.exception("Search failed")
        raise HTTPException(503, "Search isn't available right now. Please try again.")


# ---------------------------------------------------------------------------
# "Recommended for you": the Hybrid recommendation system (content-based
# filtering + collaborative filtering + ranking, see recommendations.py)
# ---------------------------------------------------------------------------

@app.get("/recommendations")
def get_recommendations(authorization: str | None = Header(default=None)):
    """The dashboard home's "Recommended for you": job posts for freelancers,
    services for clients, best first, with the reasons they were picked.
    Returns ids only; the website loads the posts itself, under the normal
    database rules."""
    user_id = get_user_id(authorization)
    try:
        return recommend(supabase, user_id)
    except APIError:
        logging.exception("Recommendations failed")
        raise HTTPException(503, "Recommendations aren't available right now.")


# ---------------------------------------------------------------------------
# AI Moodboard Matching: feature 2 in PhilFreela-System-Functions.md
# (CLIP embeddings + cosine similarity, see moodboard.py)
# ---------------------------------------------------------------------------

@app.post("/moodboard/match")
def moodboard_match(image: UploadFile = File(...), authorization: str | None = Header(default=None)):
    """A client uploads a reference image (a moodboard, a sample, a style
    they like); returns the freelancers whose portfolio work looks closest
    to it, best first. Returns ids only; the website loads the freelancer
    and the matching slide itself, under the normal database rules."""
    get_user_id(authorization)
    reference = read_image(image, "reference image")
    try:
        return {"results": match_moodboard(supabase, reference)}
    except APIError:
        logging.exception("Moodboard matching failed")
        raise HTTPException(503, "Moodboard matching isn't available right now. Please try again.")
